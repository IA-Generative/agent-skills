#!/usr/bin/env python3
"""Rend un document d'homologation MirAI au format ministériel `.docx`.

Construit le `.docx` **sur le modèle** `MODELE_DAT_MirAI.docx` (hérite page de garde,
styles, sommaire/TOC) à partir d'un contenu **Markdown fidèle à la structure du modèle**
(tel que produit par la skill `homologation-dat`).

Opérations :
  - conserve la page de garde et le champ Sommaire (TOC) du modèle ;
  - supprime la page « Comment utiliser ce modèle » et tous les encadrés d'instruction (« ✎ ») ;
  - supprime le squelette de corps du modèle et le remplace par le contenu fourni ;
  - applique les styles de titre RÉELS du modèle (styleId Heading1/2/3 — l'accès par nom peut échouer) ;
  - saut de page avant chaque section de niveau 1 ; `keep_with_next` sur les titres ;
  - tableaux Markdown → tableaux Word (en-tête répété `tblHeader`, lignes non sécables `cantSplit`) ;
  - blocs de code ``` → police à chasse fixe (rend l'ASCII des schémas) ;
  - marque le TOC « dirty » + `settings/updateFields=true` (sinon sommaire vide ;
    repli utilisateur : Ctrl+A puis F9).

Mapping des niveaux Markdown → modèle : `##` → Heading1, `###` → Heading2, `####` → Heading3.
Le titre `#` et le bloc de méta en tête (avant le premier `##`) sont ignorés (déjà sur la garde).

Usage :
    python render_docx.py --input DAT.md --output private/DAT_x_v0.1.docx \
        [--model /chemin/MODELE_DAT_MirAI.docx] [--service "Nom"] [--version v0.1] [--date "21 juin 2026"]

Dépendance : python-docx (`pip install python-docx`).
"""
from __future__ import annotations

import argparse
import os
import re
import sys

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.exit("python-docx requis : pip install python-docx")

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
LEVEL_STYLE = {2: "Heading1", 3: "Heading2", 4: "Heading3"}  # nb de '#' -> styleId modèle


# ── Helpers XML sur le corps du modèle ────────────────────────────────────────

def p_text(el) -> str:
    return "".join(t.text or "" for t in el.findall(".//" + qn("w:t")))


def has_toc(el) -> bool:
    """True si l'élément (w:p ou w:sdt) contient un champ TOC dans ses descendants."""
    return any("TOC" in (it.text or "") for it in el.iter(qn("w:instrText")))


def strip_template_body(doc) -> None:
    """Retire la page how-to et le squelette ; conserve garde + Sommaire + TOC + sectPr.

    Robuste au TOC encapsulé dans un w:sdt (Word moderne) : les repères sont cherchés
    parmi les enfants DIRECTS du corps via leurs descendants textuels.
    """
    body = doc.element.body
    children = list(body)

    def find(pred):
        # cherche un enfant direct dont un descendant w:p correspond au prédicat
        for el in children:
            for p in ([el] if el.tag == qn("w:p") else el.iter(qn("w:p"))):
                if pred(p_text(p).strip()):
                    return el
        return None

    howto = find(lambda t: t.startswith("Comment utiliser ce mod"))
    sommaire = find(lambda t: t == "Sommaire")
    toc = next((el for el in children if has_toc(el)), None)
    sectpr = body.find(qn("w:sectPr"))

    if howto is None or sommaire is None or toc is None:
        raise SystemExit("Modèle inattendu : repères (how-to / Sommaire / TOC) introuvables.")

    i_howto, i_somm, i_toc = children.index(howto), children.index(sommaire), children.index(toc)

    # 1) supprime la page how-to (de 'Comment utiliser' jusqu'avant 'Sommaire')
    for el in children[i_howto:i_somm]:
        body.remove(el)
    # 2) supprime le squelette (après le TOC jusqu'au sectPr)
    for el in children[i_toc + 1:]:
        if el is not sectpr:
            body.remove(el)
    # 3) nettoie les paragraphes vides en fin de page de garde (évite une page blanche)
    cover = [el for el in list(body)[: list(body).index(sommaire)] if el.tag == qn("w:p")]
    for el in reversed(cover):
        if p_text(el).strip():
            break
        body.remove(el)

    # saut de page avant le Sommaire pour le détacher proprement de la garde
    _set_page_break_before(sommaire)


def mark_toc_dirty(doc) -> None:
    for el in doc.element.body.iter(qn("w:fldChar")):
        if el.get(qn("w:fldCharType")) == "begin":
            el.set(qn("w:dirty"), "true")
    # settings/updateFields = true
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)


def fill_cover(doc, service, version, date) -> None:
    """Remplace les libellés de la page de garde (par run, sur tout le doc)."""
    repl = {"— MODÈLE": "", "— MODELE": "", "- MODÈLE": ""}
    if service:
        repl["[NOM DU SERVICE]"] = service
    if date:
        repl["[JJ mois AAAA]"] = date
    for para in doc.paragraphs:
        full = "".join(r.text for r in para.runs)
        new = full
        for k, v in repl.items():
            new = new.replace(k, v)
        if version and full.strip().startswith("Version"):
            new = re.sub(r"Version\s*:.*", f"Version : {version}", new)
        if new != full and para.runs:
            for r in para.runs[1:]:
                r.text = ""
            para.runs[0].text = new


# ── Styles & mise en forme ────────────────────────────────────────────────────

def _set_style_id(paragraph, style_id) -> None:
    """Applique un style par son styleId réel (robuste vs accès par nom)."""
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:pStyle")):
        pPr.remove(old)
    pStyle = OxmlElement("w:pStyle")
    pStyle.set(qn("w:val"), style_id)
    pPr.insert(0, pStyle)


def _set_page_break_before(el) -> None:
    pPr = el.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        el.insert(0, pPr)
    if pPr.find(qn("w:pageBreakBefore")) is None:
        pPr.append(OxmlElement("w:pageBreakBefore"))


def _add_runs(paragraph, text) -> None:
    """Rend le gras `**...**` et retire les backticks `code`."""
    text = text.replace("`", "")
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        if i % 2 == 1:
            run.bold = True


def _set_table_borders(table) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "999999")
        borders.append(e)
    tblPr.append(borders)


def _add_table(doc, rows):
    header, body_rows = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    try:
        table.style = "Table Grid"
    except Exception:
        _set_table_borders(table)
    for j, cell in enumerate(header):
        para = table.rows[0].cells[j].paragraphs[0]
        _add_runs(para, cell)
        for r in para.runs:
            r.bold = True
    # en-tête répété + ligne non sécable
    trPr = table.rows[0]._tr.get_or_add_trPr()
    for tag in ("w:tblHeader", "w:cantSplit"):
        trPr.append(OxmlElement(tag))
    for row in body_rows:
        cells = table.add_row().cells
        tr = table.rows[-1]._tr
        tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, val in enumerate(row):
            if j < len(cells):
                _add_runs(cells[j].paragraphs[0], val)
    return table


# ── Parsing Markdown → corps Word ─────────────────────────────────────────────

def _split_table_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def render_markdown(doc, md) -> None:
    lines = md.splitlines()
    started = False  # on commence au premier '## '
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        h = re.match(r"^(#{1,6})\s+(.*)$", line)
        if h and len(h.group(1)) >= 2:
            started = True
        if not started:
            i += 1
            continue

        # Titres de section
        if h:
            level = len(h.group(1))
            style = LEVEL_STYLE.get(level)
            if style:
                p = doc.add_paragraph()
                _set_style_id(p, style)
                _add_runs(p, h.group(2).strip())
                p.paragraph_format.keep_with_next = True
                if style == "Heading1":
                    _set_page_break_before(p._p)
            i += 1
            continue

        # Bloc de code ``` -> police à chasse fixe
        if line.strip().startswith("```"):
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                run = p.add_run(lines[i])
                run.font.name = "Consolas"
                i += 1
            i += 1
            continue

        # Tableau Markdown
        if line.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|?[\s:|-]+\|", lines[i + 1]):
            rows = [_split_table_row(line)]
            i += 2  # saute le séparateur
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append(_split_table_row(lines[i]))
                i += 1
            _add_table(doc, rows)
            doc.add_paragraph()
            continue

        # Liste à puces (style ListParagraph du modèle + glyphe •)
        mli = re.match(r"^\s*[-*]\s+(.*)$", line)
        if mli:
            p = doc.add_paragraph()
            _set_style_id(p, "ListParagraph")
            _add_runs(p, "•  " + mli.group(1))
            i += 1
            continue

        # Citation / encadré
        mq = re.match(r"^>\s?(.*)$", line)
        if mq:
            p = doc.add_paragraph()
            _add_runs(p, mq.group(1))
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # Règle horizontale -> ignorée
        if re.match(r"^\s*---+\s*$", line):
            i += 1
            continue

        # Paragraphe normal
        if line.strip():
            p = doc.add_paragraph()
            _add_runs(p, line.strip())
        i += 1


def reorder_sectpr(doc) -> None:
    body = doc.element.body
    sectpr = body.find(qn("w:sectPr"))
    if sectpr is not None:
        body.remove(sectpr)
        body.append(sectpr)


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    ap = argparse.ArgumentParser(description="Rend un DAT MirAI en .docx sur le modèle ministériel.")
    ap.add_argument("--input", required=True, help="Contenu Markdown fidèle à la structure du modèle")
    ap.add_argument("--output", required=True, help="Chemin du .docx de sortie")
    here = os.path.dirname(os.path.abspath(__file__))
    default_model = os.path.normpath(os.path.join(here, "..", "homologation-dat", "dat", "MODELE_DAT_MirAI.docx"))
    ap.add_argument("--model", default=default_model, help="Modèle .docx (défaut : skill homologation-dat)")
    ap.add_argument("--service", default="", help="Nom du service (page de garde)")
    ap.add_argument("--version", default="", help="Version (page de garde)")
    ap.add_argument("--date", default="", help="Date (page de garde)")
    args = ap.parse_args()

    if not os.path.isfile(args.model):
        sys.exit(f"Modèle introuvable : {args.model}")
    with open(args.input, encoding="utf-8") as f:
        md = f.read()

    doc = Document(args.model)
    strip_template_body(doc)
    fill_cover(doc, args.service, args.version, args.date)
    render_markdown(doc, md)
    reorder_sectpr(doc)
    mark_toc_dirty(doc)

    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    doc.save(args.output)
    print(f"OK → {args.output}")


if __name__ == "__main__":
    main()
